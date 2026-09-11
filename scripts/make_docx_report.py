"""Generate an editable .docx version of the summary report -- real Word tables,
full explanations, easy to open in Word/LibreOffice/Google Docs and edit directly."""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = "results/summary_report.docx"

INK = RGBColor(0x14, 0x18, 0x1B)
DIM = RGBColor(0x4A, 0x52, 0x59)
ACCENT = RGBColor(0x2E, 0x6F, 0x95)
HEADER_BG = "2E6F95"
HILITE_BG = "FBF1E4"


def set_cell_shading(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def style_cell(cell, text, bold=False, color=None, size=10, align_center=False):
    cell.text = ""
    p = cell.paragraphs[0]
    if align_center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = "Calibri"
    if color:
        run.font.color.rgb = color


def add_table(doc, headers, rows, col_widths_cm, highlight_row_idxs=()):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    t.autofit = False

    for j, h in enumerate(headers):
        cell = t.rows[0].cells[j]
        style_cell(cell, h, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF), size=10,
                   align_center=(j > 0))
        set_cell_shading(cell, HEADER_BG)

    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = t.rows[i + 1].cells[j]
            style_cell(cell, val, bold=False, color=INK, size=10, align_center=(j > 0))
            if i in highlight_row_idxs:
                set_cell_shading(cell, HILITE_BG)

    for row in t.rows:
        for j, w in enumerate(col_widths_cm):
            row.cells[j].width = Cm(w)
    doc.add_paragraph()
    return t


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    _add_runs(p, text, size=9, color=DIM, italic=True)


def add_heading(doc, text, size=15):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(size)
    run.font.color.rgb = INK
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    return p


def add_body(doc, text, size=11):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    _add_runs(p, text, size)
    return p


def _add_runs(p, text, size, color=None, italic=False):
    """Minimal **bold** / *italic* markdown-ish parsing so callers can write inline emphasis."""
    import re
    parts = re.split(r"(\*\*.*?\*\*|\*.*?\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = p.add_run(part[2:-2])
            r.bold = True
            r.italic = italic
        elif part.startswith("*") and part.endswith("*"):
            r = p.add_run(part[1:-1])
            r.italic = True
        else:
            r = p.add_run(part)
            r.italic = italic
        r.font.size = Pt(size)
        r.font.color.rgb = color or INK


def add_bullets(doc, items, size=10.5):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(6)
        _add_runs(p, item, size)


doc = Document()

# base font
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)
for section in doc.sections:
    section.left_margin = Cm(2)
    section.right_margin = Cm(2)
    section.top_margin = Cm(1.7)
    section.bottom_margin = Cm(1.7)

# ---------- Title ----------
p = doc.add_paragraph()
run = p.add_run("INT8 Detection Under Degradation")
run.bold = True
run.font.size = Pt(22)
run.font.color.rgb = INK
p.paragraph_format.space_after = Pt(2)

p = doc.add_paragraph()
run = p.add_run("YOLOv8n on COCO val2017 — 5 classes, 500-image eval set — "
                "12th Gen Intel Core i5-1245U, single-threaded")
run.font.size = Pt(11)
run.font.color.rgb = DIM
p.paragraph_format.space_after = Pt(16)

# ---------- What we did ----------
add_heading(doc, "What we did")
add_body(doc,
    "We picked **YOLOv8n** because it runs on CPU with no special hardware, ships with pretrained "
    "COCO weights, and has a mature ONNX/INT8 export path — the tooling maturity mattered most "
    "since quantization was the whole point of this project."
)
add_body(doc,
    "We built a fixed **500-image subset of COCO val2017**, restricted to 5 classes "
    "(person, car, bicycle, traffic light, stop sign) — 2,353 ground-truth boxes in total. "
    "A separate, disjoint pool of 120 images was set aside purely for INT8 calibration, so calibration "
    "data never overlaps with what we evaluate on."
)
add_body(doc,
    "**Task 1 — FP32 baseline, then quantize to INT8.** Exported the pretrained model to a fixed "
    "640×640 ONNX graph and measured accuracy with pycocotools.COCOeval. Then applied post-training "
    "static INT8 quantization via ONNX Runtime (QDQ format, U8S8, per-channel weights, Conv-only — "
    "quantizing the detection head's final output layer destroys every class score, so it stays in "
    "float32), calibrated on the 120-image set. Compared accuracy, latency, and model size against FP32."
)
add_body(doc,
    "**Task 2 — Degrade.** Built four degraded copies of the same 500 images (motion blur, low light, "
    "JPEG compression, downscale/upscale) and ran both FP32 and INT8 across all four, to see whether "
    "quantization's accuracy cost gets worse under realistic image degradation."
)
add_body(doc,
    "**Task 3 — one targeted intervention.** Picked the worst-performing degradation (motion blur) "
    "and tried a single, targeted fix — recalibrating INT8's quantization scales on blurred images "
    "instead of clean ones — to test whether that closed the gap."
)

# ---------- Results ----------
add_heading(doc, "Results")

add_body(doc, "**Table 1. FP32 vs. INT8 — accuracy, latency, and model size (clean images).**", size=11)
add_table(doc,
    headers=["Metric", "FP32", "INT8", "Delta"],
    rows=[
        ["mAP@0.5", "0.586", "0.569", "−3.0%"],
        ["mAP@0.5:0.95", "0.415", "0.394", "−4.9%"],
        ["Latency, mean (1 thread)", "134.9 ms", "69.7 ms", "1.94x faster"],
        ["Latency, p95 (1 thread)", "138.6 ms", "72.0 ms", "1.93x faster"],
        ["Model size (ONNX)", "12.85 MB", "3.60 MB", "3.57x smaller"],
    ],
    col_widths_cm=[5.5, 3, 3, 3.5],
)
add_caption(doc,
    "INT8 gives roughly a 2x speedup and 3.6x size reduction for a 3–5% relative accuracy cost. "
    "Small-object AP took the largest hit (−18.9% relative) while large-object AP was essentially "
    "unaffected — 8-bit resolution has the least room to spare exactly where the signal is thinnest.")

add_body(doc, "**Table 2. mAP@0.5 across all five conditions — the 5×2 table — FP32 vs. INT8.**", size=11)
add_table(doc,
    headers=["Condition", "FP32", "INT8", "INT8 relative gap"],
    rows=[
        ["Low light (γ = 2.5 gamma correction)", "0.516", "0.512", "−0.6%"],
        ["Clean (no degradation)", "0.586", "0.569", "−3.0%"],
        ["JPEG compression, quality 30", "0.534", "0.514", "−3.6%"],
        ["Downscale to 50%, then upscale back", "0.543", "0.517", "−4.8%"],
        ["Motion blur (15×15 horizontal kernel)", "0.281", "0.250", "−10.9%"],
    ],
    col_widths_cm=[6.5, 3, 3, 3.5],
    highlight_row_idxs=(4,),
)
add_caption(doc,
    "“INT8 relative gap” is INT8's shortfall as a fraction of FP32's own score in that condition, "
    "so it's comparable across rows even though FP32 itself collapses under motion blur. Reading the "
    "table: INT8's accuracy penalty relative to FP32 does **not** grow uniformly with how visually "
    "damaging the degradation looks — it barely moves under low light, but nearly quadruples under "
    "motion blur, which is also the single worst condition for both models by far.")

add_body(doc, "**Table 3. Task 3 intervention — recalibrating INT8 on blurred images.**", size=11)
add_table(doc,
    headers=["Model", "mAP@0.5 on motion blur", "mAP@0.5 on clean images"],
    rows=[
        ["FP32", "0.281", "0.586"],
        ["INT8, clean-calibrated (original)", "0.250", "0.569"],
        ["INT8, blur-calibrated (new)", "0.250", "0.582"],
    ],
    col_widths_cm=[6.5, 5, 5],
    highlight_row_idxs=(2,),
)
add_caption(doc,
    "The intervention **did not** close the gap on its intended target (motion blur: 0.250 → 0.250, "
    "no real change). It did, unexpectedly, close 78% of the original gap on clean images instead "
    "(0.569 → 0.582) — recalibrating on blurred data made 95% of the model's 124 quantization scales "
    "tighter, which improved resolution for the clean distribution's typical activation values but did "
    "nothing for information already destroyed by the blur itself. Cost: +6 bytes, <0.3% latency "
    "difference — essentially free, just ineffective for the stated goal.")

# ---------- What surprised us ----------
add_heading(doc, "What surprised us")
add_bullets(doc, [
    "The very first INT8 build was **slower** than FP32 (177 ms vs. 135 ms per image). The cause: "
    "we'd used S8S8 quantization (signed int8 for both activations and weights), and ONNX Runtime's "
    "accelerated x86-64 integer kernels are built for **U8S8**, not S8S8 — S8S8 silently falls back "
    "to a slower, unaccelerated code path. Switching one parameter (activation dtype) got the expected "
    "~2x speedup, on the exact same CPU.",
    "Quantizing the model's own output layer — a Concat that merges box coordinates (range ~0–640) "
    "with class-confidence scores (range 0–1) into one tensor — gives that whole tensor one shared "
    "quantization scale. Calibrated for the 0–640 range, it rounds every class score below about 2.5 "
    "down to exactly zero. Every detection silently disappeared, with no error or crash.",
    "INT8's accuracy penalty relative to FP32 does not scale with how bad a degradation *looks*. Low "
    "light is visually dramatic (mean pixel brightness drops ~60%) but barely widens INT8's gap. Motion "
    "blur, in contrast, nearly quadruples it — because blur specifically destroys the fine edge "
    "information that INT8's already-reduced numeric resolution depends on most.",
    "Trying to fix the motion-blur gap by recalibrating on blurred images didn't fix motion blur at "
    "all — it fixed the clean-image gap instead, a case we weren't even targeting. The mechanism: "
    "recalibration changes *where* the fixed 8-bit precision budget gets spent, not how much "
    "information a destructive degradation has already removed.",
])

# ---------- Next week ----------
add_heading(doc, "What we'd do with another week")
add_bullets(doc, [
    "**Mixed precision instead of a calibration swap.** Keep the earliest backbone layers "
    "(model.1 / model.2 — the layers that showed the largest quantization-scale shrinkage under "
    "blur-calibration) in FP32, since that's exactly where blur destroys the edge/gradient information "
    "those layers extract. This targets the actual mechanism instead of reshuffling quantization "
    "resolution after the fact.",
    "**Quantization-aware training (QAT) with motion blur in the augmentation pipeline**, so the "
    "network's weights themselves adapt to the degradation, rather than only recalibrating the "
    "quantization ranges around fixed, already-trained weights.",
    "**Widen the eval set** to the full 5-class COCO val2017 pool (roughly 3,000 images). Some of our "
    "per-class and per-size splits are thin right now — `stop sign` has only 15 ground-truth instances "
    "total — so those specific numbers carry real statistical noise.",
    "**Test more degradations that plausibly compound with blur in real deployment** (sensor noise, "
    "genuinely low-resolution capture, weather effects), and benchmark **OpenVINO's INT8 path** on the "
    "same CPU as a second quantization backend, since it's purpose-built for Intel hardware and might "
    "handle the motion-blur case differently than ONNX Runtime's generic CPU kernels.",
])

doc.save(OUT)
print("Wrote", OUT)
