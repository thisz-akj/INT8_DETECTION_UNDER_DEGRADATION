"""One-time surgical patch: fill in the blank spacer paragraph right after the
per-class AP table with a short explanatory caption, matching the style of the
other table captions in results/summary_report.docx."""
from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn

PATH = "results/summary_report.docx"
DIM = RGBColor(0x4A, 0x52, 0x59)

CAPTION = (
    "person and stop sign score highest, bicycle and traffic light lowest — largely "
    "tracking object size and how visually distinctive each class is; stop sign's "
    "numbers are also the noisiest here, with only 15 ground-truth instances in the eval set."
)

doc = Document(PATH)

# find the "Table 2. Per-class AP..." heading paragraph
heading_p = None
for p in doc.paragraphs:
    if p.text.strip().startswith("Table 2. Per-class AP"):
        heading_p = p
        break
if heading_p is None:
    raise SystemExit("Could not find the per-class AP table heading -- aborting, nothing saved.")

# structurally: heading -> table -> spacer paragraph (currently empty)
table_el = heading_p._p.getnext()
if table_el is None or table_el.tag != qn("w:tbl"):
    raise SystemExit("Expected a table right after the heading -- aborting, nothing saved.")
spacer_el = table_el.getnext()
if spacer_el is None or spacer_el.tag != qn("w:p"):
    raise SystemExit("Expected a spacer paragraph right after the table -- aborting, nothing saved.")

from docx.text.paragraph import Paragraph
spacer = Paragraph(spacer_el, heading_p._parent)
if spacer.text.strip():
    raise SystemExit("Spacer paragraph isn't empty as expected -- aborting, nothing saved.")

run = spacer.add_run(CAPTION)
run.italic = True
run.font.size = Pt(9)
run.font.color.rgb = DIM

doc.save(PATH)
print("Patched", PATH)
